# -*- coding: utf-8 -*-
"""
甍AI Estimation OS — Geometry Engine + Quantity View + Document Engine
（e0.2 試作 / Data Model Complete）

思想（DATAMODEL.md より）:
    Project → Decision → Measurement(vertices) → Geometry Engine
             → Quantity View → Estimate → Document Engine → Template

このスクリプトは e0.2 の「証明の種」。真実は Measurement.vertices だけ。
  - amount は保存しない（Geometry Engine が vertices から計算）。
  - Quantity は保存しない（Document Engine が Measurement を集約して毎回生成する派生ビュー）。
  - Geometry Engine は拡張しない。計算能力は GEOMETRY_KNOWLEDGE（辞典）が持つ。
本格実装は e0.4 Geometry Engine / e0.5 Document Engine。ここは最小の種。
"""

import json
import sys
import math
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FONT = "Yu Gothic"

THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="2E74B5")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=FONT, bold=True, size=16, color="1F4E79")
LABEL_FONT = Font(name=FONT, bold=True, size=10, color="1F4E79")
BODY_FONT = Font(name=FONT, size=10)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


# ======================================================================
#  Geometry Knowledge（辞典）: operation → calculator
#  新しい計算(Perimeter/Centroid…)が増えても、ここに1行足すだけ。
#  Geometry Engine 本体は一切変えない（DATAMODEL 第4章・不変条件10）。
# ======================================================================
def _polygon_area(v):
    n = len(v)
    s = sum(v[i][0] * v[(i + 1) % n][1] - v[(i + 1) % n][0] * v[i][1] for i in range(n))
    return round(abs(s) / 2.0, 2)


def _polyline_length(v):
    return round(sum(math.dist(v[i], v[i + 1]) for i in range(len(v) - 1)), 2)


def _point_count(v):
    return len(v)


GEOMETRY_KNOWLEDGE = {
    "Area":   _polygon_area,      # Polygon → Area
    "Length": _polyline_length,   # Line / Polyline → Length
    "Count":  _point_count,       # Point → Count
    # 将来: "Perimeter": _polygon_perimeter, "Centroid": _polygon_centroid …
}


# ======================================================================
#  Geometry Engine: (geometry, operation) を受け取り、Knowledge に委譲するだけ
# ======================================================================
def geometry_calculate(m):
    op = m["operation"]
    calc = GEOMETRY_KNOWLEDGE.get(op)
    if calc is None:
        raise ValueError(f"Geometry Knowledge に operation '{op}' が無い")
    return calc(m["vertices"])


# ======================================================================
#  Quantity View: Measurement を (trade,item,unit) で集約する派生ビュー
#  保存しない。呼ばれるたびに Measurement から生成する。
# ======================================================================
def quantity_view(data):
    groups = OrderedDict()
    for m in data["measurements"]:
        key = (m["trade"], m["item"], m["unit"])
        g = groups.setdefault(key, {"trade": key[0], "item": key[1], "unit": key[2],
                                    "measurements": [], "labels": [], "amount": 0.0})
        g["measurements"].append(m["measurementId"])
        g["labels"].append(m["label"])
        g["amount"] = round(g["amount"] + geometry_calculate(m), 2)
    for g in groups.values():
        if g["amount"] == int(g["amount"]):
            g["amount"] = int(g["amount"])  # 個数や整長は整数表示
    return list(groups.values())


# ---- データ層 -------------------------------------------------------------
def load_data(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def check_invariants(data):
    """DATAMODEL.md 第9章 不変条件を検査する。破れば帳票は出さない。"""
    dids = {d["decisionId"] for d in data["decisions"]}
    errors = []

    if "quantities" in data:
        errors.append("quantities を保存している（Quantity は派生ビュー・不変条件5違反）")

    for d in data["decisions"]:
        if not d.get("basis"):
            errors.append(f"{d['decisionId']}: basis が空（原則13違反）")

    for m in data["measurements"]:
        if "amount" in m:
            errors.append(f"{m['measurementId']}: amount を保存している（不変条件1違反）")
        if not m.get("vertices"):
            errors.append(f"{m['measurementId']}: vertices が空（拾いの真実がない）")
        for f in ("trade", "item", "unit"):
            if not m.get(f):
                errors.append(f"{m['measurementId']}: {f} が空（集約先が定まらない）")
        if not m.get("basisDecision"):
            errors.append(f"{m['measurementId']}: basisDecision が空")
        for ref in m.get("basisDecision", []):
            if ref not in dids:
                errors.append(f"{m['measurementId']}: 根拠 {ref} が存在しない Decision")
        if m["operation"] not in GEOMETRY_KNOWLEDGE:
            errors.append(f"{m['measurementId']}: operation '{m['operation']}' が Geometry Knowledge に無い")

    return errors


# ---- テンプレート層 -------------------------------------------------------
def _header(ws, row, headers, widths):
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill, c.font, c.alignment, c.border = HEAD_FILL, HEAD_FONT, CENTER, BORDER
        ws.column_dimensions[c.column_letter].width = w


def _rows(ws, start, records, aligns):
    r = start
    for rec in records:
        for col, (v, a) in enumerate(zip(rec, aligns), start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.font, c.alignment, c.border = BODY_FONT, a, BORDER
        r += 1
    return r


def render_sekisan(ws, data):
    """Document(1): 積算表。Quantity View（毎回生成）から描く。"""
    proj = data["project"]
    ws.title = "積算表"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:F1")
    ws["A1"] = "積 算 表"
    ws["A1"].font, ws["A1"].alignment = TITLE_FONT, CENTER
    ws.row_dimensions[1].height = 28
    ws["A2"], ws["A2"].font = "案件", LABEL_FONT
    ws.merge_cells("B2:D2")
    ws["B2"], ws["B2"].font = proj["name"], BODY_FONT
    ws["E2"], ws["E2"].font = "積算ID", LABEL_FONT
    ws["F2"], ws["F2"].font = proj["extensions"]["estimation"]["estimationId"], BODY_FONT

    _header(ws, 4, ["工種", "名称", "数量", "単位", "根拠(拾い)", "内訳"], [12, 22, 9, 7, 20, 26])
    recs = [[g["trade"], g["item"], g["amount"], g["unit"],
             " / ".join(g["measurements"]), " ＋ ".join(g["labels"])] for g in quantity_view(data)]
    end = _rows(ws, 5, recs, [CENTER, LEFT, RIGHT, CENTER, CENTER, LEFT])

    note = ("※ 数量は Geometry Engine が vertices から計算し、Quantity View が集約した派生値（保存ゼロ）。"
            "積算表 → 拾い明細 → 判断台帳 と辿れる。金額は含まない。")
    ws.merge_cells(start_row=end + 1, start_column=1, end_row=end + 1, end_column=6)
    c = ws.cell(row=end + 1, column=1, value=note)
    c.font, c.alignment = Font(name=FONT, size=9, italic=True, color="666666"), LEFT


def render_measurement_log(ws, data):
    """拾い明細: geometry と計算値（Quantity View → Measurement）。"""
    ws.title = "拾い明細"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    ws["A1"] = "拾い明細（Measurement / geometry → 計算）"
    ws["A1"].font, ws["A1"].alignment = TITLE_FONT, CENTER
    ws.row_dimensions[1].height = 28

    _header(ws, 3, ["ID", "operation", "geometry", "拾い", "計算値", "単位", "集約先(item)", "根拠(判断)"],
            [9, 9, 9, 12, 8, 6, 16, 16])
    recs = [[m["measurementId"], m["operation"], m["geometry"], m["label"],
             geometry_calculate(m), m["unit"], m["item"], " / ".join(m["basisDecision"])]
            for m in data["measurements"]]
    _rows(ws, 4, recs, [CENTER, CENTER, CENTER, LEFT, RIGHT, CENTER, LEFT, CENTER])


def render_decision_log(ws, data):
    """判断台帳: 拾いの『根拠(判断)』を解決（Measurement→Decision）。"""
    ws.title = "判断台帳"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = "判断台帳（Decision Log）"
    ws["A1"].font, ws["A1"].alignment = TITLE_FONT, CENTER
    ws.row_dimensions[1].height = 28

    _header(ws, 3, ["ID", "種別", "判断", "根拠(basis)", "確信度", "作成者"], [9, 16, 14, 30, 9, 10])
    recs = [[d["decisionId"], d["type"], d["value"], " / ".join(d["basis"]),
             f'{d["confidence"]}%', "人" if d["author"] == "human" else "AI"]
            for d in data["decisions"]]
    _rows(ws, 4, recs, [CENTER, CENTER, CENTER, LEFT, CENTER, CENTER])


# ---- Document Engine 本体 -------------------------------------------------
DOCUMENTS = {
    "積算表": render_sekisan,
    "拾い明細": render_measurement_log,
    "判断台帳": render_decision_log,
}


def generate(data, document_names, out_path):
    wb = Workbook()
    wb.remove(wb.active)
    for name in document_names:
        DOCUMENTS[name](wb.create_sheet(), data)
    wb.save(out_path)
    return out_path


def main():
    base = Path(__file__).parent
    data = load_data(base / "sample_data.json")

    errs = check_invariants(data)
    if errs:
        print("不変条件エラー（帳票を出しません）:")
        for e in errs:
            print("  -", e)
        sys.exit(1)

    out = base / "積算表.xlsx"
    generate(data, ["積算表", "拾い明細", "判断台帳"], out)
    print("OK:", out)
    print(f"  Decision {len(data['decisions'])} / Measurement {len(data['measurements'])}"
          f" → Quantity View {len(quantity_view(data))} 行（保存ゼロ・毎回生成）")
    for g in quantity_view(data):
        print(f"    {g['item']}: {g['amount']}{g['unit']}  ← Σ{g['measurements']}")


if __name__ == "__main__":
    main()
